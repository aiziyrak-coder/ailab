#!/usr/bin/env python3
"""Yo'llanma namunasi — chop etish uchun PDF va tahrirlash uchun DOCX.

Nima uchun: klinikalar hozir turli shakldagi, qo'lda to'ldirilgan
yo'llanmalarni suratga olib yuklaydi. Shakl har xil bo'lgani uchun OCR
maydonlarni tushirib qoldiradi (masalan «Ангиома?» klinik tashxisi
yo'qolgan edi). Bitta chiroyli, aniq kataklari bo'lgan standart namuna:
shifokor to'ldiradi → laborant suratga oladi → DermaPATH o'qiydi.

Dizayn: A4, Radeski/DermaPATH brendi (tilla + qora), keng kataklar, katta
yorliqlar — telefon kamerasi ham, OCR ham bemalol o'qiydi.

    python scripts/build_referral_form.py
→ frontend/static/docs/yollanma.pdf, yollanma.docx
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "frontend" / "static" / "docs"
LOGO = ROOT / "frontend" / "static" / "logo.png"

GOLD = "#c09949"
GOLD_DEEP = "#a8813a"
GOLD_SOFT = "#e8dcc2"
GOLD_WASH = "#faf6ec"
INK = "#1a1a1a"
INK_SOFT = "#5f6664"
LINE = "#d9d4c9"

CLINIC = "RADESKI SKIN CLINIC"
PRODUCT = "DermaPATH"
FORM_NO = "014-raqamli tibbiy hujjat shakli"
TITLE = "PATOMORFOLOGIK TEKSHIRUVGA YO‘LLANMA"

# (yorliq, izoh, balandlik mm) — bo'sh bo'lsa standart balandlik
PATIENT_FIELDS = [
    ("Bemor F.I.Sh.", "familiya, ism, otasining ismi — to‘liq", 10),
    ("Tug‘ilgan sana / yoshi", "KK.OO.YYYY yoki yosh", 8),
    ("Manzil", "viloyat, tuman/shahar", 8),
    ("Telefon", "+998 __ ___ __ __", 8),
]
CLINICAL_FIELDS = [
    ("KLINIK TASHXIS", "shifokor taxmini — kasallik nomi; bir nechta bo‘lsa vergul bilan", 13),
    ("Namuna olingan joy", "organ va anatomik joy: masalan «teri, bo‘yin, o‘ng yon»", 8),
    ("Amaliyot turi va sanasi", "eksizion / insizion / panch / shave / kyuretaj — sana", 9),
    ("Klinik ma’lumot", "davomiyligi, kechishi, o‘lchami, avvalgi davolash, qo‘shimcha kasalliklar", 15),
]
LAB_FIELDS = [
    ("Gistologik №", "", 8),
    ("Gistologik xulosa", "", 16),
    ("Patologoanatom", "F.I.Sh., imzo", 8),
    ("Laborant", "F.I.Sh., imzo", 8),
]


def _font(pdfmetrics, TTFont):
    """Kirill va lotinni to‘liq qoplaydigan shrift: Arial (Windows) yoki DejaVu."""
    cands = [
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
         "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ]
    for reg, bold in cands:
        if os.path.isfile(reg) and os.path.isfile(bold):
            pdfmetrics.registerFont(TTFont("Body", reg))
            pdfmetrics.registerFont(TTFont("Body-Bold", bold))
            return "Body", "Body-Bold"
    return "Helvetica", "Helvetica-Bold"


def build_pdf(path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    REG, BOLD = _font(pdfmetrics, TTFont)
    W, H = A4
    M = 14 * mm
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle("Patomorfologik tekshiruvga yo‘llanma — DermaPATH")
    c.setAuthor(CLINIC)

    gold, gold_deep, ink, soft, line = (colors.HexColor(x) for x in (GOLD, GOLD_DEEP, INK, INK_SOFT, LINE))
    wash = colors.HexColor(GOLD_WASH)

    # ── Sarlavha ────────────────────────────────────────────────────────────
    y = H - M
    c.setFillColor(gold)
    c.rect(0, H - 6 * mm, W, 6 * mm, stroke=0, fill=1)          # tepa tilla chiziq
    if LOGO.is_file():
        try:
            c.drawImage(str(LOGO), M, y - 16 * mm, width=16 * mm, height=16 * mm,
                        preserveAspectRatio=True, mask="auto")
        except Exception:
            pass
    c.setFillColor(ink)
    c.setFont(BOLD, 15)
    c.drawString(M + 20 * mm, y - 6 * mm, CLINIC)
    c.setFont(REG, 8.5)
    c.setFillColor(soft)
    c.drawString(M + 20 * mm, y - 11 * mm, f"{PRODUCT} · gistopatologiya platformasi · lab.fermi.uz")
    c.drawString(M + 20 * mm, y - 15 * mm, "O‘zbekiston Respublikasi Sog‘liqni saqlash vazirligi · " + FORM_NO)

    # № va sana kataklari (o‘ng yuqori)
    bx_w, bx_h = 42 * mm, 9 * mm
    bx_x = W - M - bx_w
    for i, (lab, hint) in enumerate((("Yo‘llanma №", ""), ("Sana", "KK.OO.YYYY"))):
        by = y - 4 * mm - i * (bx_h + 2.5 * mm)
        c.setStrokeColor(gold_deep); c.setLineWidth(0.8)
        c.setFillColor(colors.white)
        c.rect(bx_x, by - bx_h, bx_w, bx_h, stroke=1, fill=1)
        c.setFillColor(gold_deep); c.setFont(BOLD, 7)
        c.drawString(bx_x + 2 * mm, by - 3.2 * mm, lab.upper())
        if hint:
            c.setFillColor(soft); c.setFont(REG, 6.5)
            c.drawRightString(bx_x + bx_w - 2 * mm, by - 3.2 * mm, hint)

    y -= 22 * mm
    c.setStrokeColor(gold); c.setLineWidth(1.2)
    c.line(M, y, W - M, y)
    y -= 9 * mm
    c.setFillColor(ink); c.setFont(BOLD, 13.5)
    c.drawCentredString(W / 2, y, TITLE)
    y -= 4.5 * mm
    c.setFillColor(soft); c.setFont(REG, 7.5)
    c.drawCentredString(W / 2, y, "Iltimos, bosma harflarda, aniq va to‘liq to‘ldiring — shakl DermaPATH tomonidan avtomatik o‘qiladi")
    y -= 7 * mm

    def section(title, y):
        c.setFillColor(wash)
        c.rect(M, y - 6.5 * mm, W - 2 * M, 6.5 * mm, stroke=0, fill=1)
        c.setFillColor(gold_deep); c.setFont(BOLD, 8.5)
        c.drawString(M + 3 * mm, y - 4.6 * mm, title.upper())
        return y - 6.5 * mm - 2 * mm

    def field(label, hint, h, y, label_w=48 * mm):
        h = h * mm
        c.setStrokeColor(line); c.setLineWidth(0.6)
        c.setFillColor(colors.white)
        c.rect(M, y - h, W - 2 * M, h, stroke=1, fill=1)
        c.setStrokeColor(line)
        c.line(M + label_w, y, M + label_w, y - h)
        c.setFillColor(ink); c.setFont(BOLD, 8.5)
        c.drawString(M + 3 * mm, y - 5 * mm, label)
        if hint:
            c.setFillColor(soft); c.setFont(REG, 6.8)
            c.drawString(M + label_w + 3 * mm, y - h + 2.2 * mm, hint)
        return y - h - 1.3 * mm

    def checks(label, options, y, h=8):
        h = h * mm
        c.setStrokeColor(line); c.setLineWidth(0.6); c.setFillColor(colors.white)
        c.rect(M, y - h, W - 2 * M, h, stroke=1, fill=1)
        c.line(M + 48 * mm, y, M + 48 * mm, y - h)
        c.setFillColor(ink); c.setFont(BOLD, 8.5)
        c.drawString(M + 3 * mm, y - 5 * mm, label)
        x = M + 48 * mm + 4 * mm
        c.setFont(REG, 8.5)
        for opt in options:
            c.setStrokeColor(gold_deep); c.setLineWidth(0.8); c.setFillColor(colors.white)
            c.rect(x, y - 6.4 * mm, 3.8 * mm, 3.8 * mm, stroke=1, fill=1)
            c.setFillColor(ink)
            c.drawString(x + 5.5 * mm, y - 5.6 * mm, opt)
            x += 5.5 * mm + pdfmetrics.stringWidth(opt, REG, 8.5) + 8 * mm
        return y - h - 1.3 * mm

    # ── 1. Bemor ────────────────────────────────────────────────────────────
    y = section("1 · Bemor", y)
    y = field(*PATIENT_FIELDS[0], y)
    y = checks("Jinsi", ["Erkak", "Ayol"], y)
    for f in PATIENT_FIELDS[1:]:
        y = field(*f, y)

    # ── 2. Klinik ma'lumot ──────────────────────────────────────────────────
    y = section("2 · Klinik ma’lumot (davolovchi shifokor to‘ldiradi)", y)
    y = field(*CLINICAL_FIELDS[0], y)
    y = field(*CLINICAL_FIELDS[1], y)
    y = checks("Amaliyot turi", ["Eksizion", "Insizion", "Panch", "Shave", "Kyuretaj", "Boshqa"], y)
    y = field("Amaliyot sanasi", "KK.OO.YYYY", 8, y)
    y = field(*CLINICAL_FIELDS[3], y)
    y = checks("Shoshilinchlik", ["Oddiy", "STAT (shoshilinch)"], y)
    y = checks("Ilova", ["Klinik surat (toshma)", "Dermatoskopiya", "Avvalgi gistologiya"], y)
    y = field("Davolovchi shifokor", "F.I.Sh., telefon, imzo", 8, y)
    y = field("Muassasa", "klinika / poliklinika nomi, shahar", 8, y)

    # ── 3. Laboratoriya ─────────────────────────────────────────────────────
    y = section("3 · Patomorfologik tekshiruv (laboratoriya to‘ldiradi)", y)
    y = field("Qabul qilingan sana", "KK.OO.YYYY · qabul qilgan xodim", 8, y)
    for f in LAB_FIELDS:
        y = field(*f, y)

    # ── Pastki qism ─────────────────────────────────────────────────────────
    c.setStrokeColor(gold); c.setLineWidth(1.2)
    c.line(M, 16 * mm, W - M, 16 * mm)
    c.setFillColor(soft); c.setFont(REG, 6.8)
    c.drawString(M, 12 * mm, "Yo‘llanma va klinik surat DermaPATH ga birga yuklanadi: suratni tekis, yorug‘ joyda, butun varaq kadrda bo‘lib oling.")
    c.drawString(M, 8.5 * mm, "Yakuniy tashxis faqat litsenziyali patolog tomonidan qo‘yiladi. " + CLINIC + " · " + PRODUCT)
    c.setFillColor(gold)
    c.rect(0, 0, W, 3 * mm, stroke=0, fill=1)
    c.showPage()
    c.save()


def build_docx(path):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.4))
    sec.top_margin, sec.bottom_margin = Cm(1.2), Cm(1.2)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(9)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    def shade(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color.lstrip("#"))
        tcPr.append(shd)

    def para(text, size=9, bold=False, color=INK, align=None, space_after=2):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold; r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(0)
        return p

    para(CLINIC, 14, True, INK, space_after=0)
    para(f"{PRODUCT} · gistopatologiya platformasi · lab.fermi.uz", 8, False, INK_SOFT, space_after=0)
    para("O‘zbekiston Respublikasi Sog‘liqni saqlash vazirligi · " + FORM_NO, 8, False, INK_SOFT, space_after=6)
    t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.RIGHT
    for cell, lab in zip(t.rows[0].cells, ("Yo‘llanma №: ______________", "Sana: ____.____.________")):
        cell.text = ""; r = cell.paragraphs[0].add_run(lab); r.bold = True; r.font.size = Pt(9)
    para("", 4, space_after=2)
    para(TITLE, 13, True, INK, "center", space_after=1)
    para("Iltimos, bosma harflarda, aniq va to‘liq to‘ldiring — shakl DermaPATH tomonidan avtomatik o‘qiladi",
         7.5, False, INK_SOFT, "center", space_after=6)

    def section(title):
        t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
        cell = t.rows[0].cells[0]; shade(cell, GOLD_WASH)
        cell.text = ""; r = cell.paragraphs[0].add_run(title.upper()); r.bold = True; r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(GOLD_DEEP.lstrip("#"))
        para("", 2, space_after=1)

    def rows(fields):
        t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
        for label, hint, h in fields:
            row = t.add_row()
            a, b = row.cells
            a.width, b.width = Cm(4.6), Cm(13.6)
            a.text = ""; r = a.paragraphs[0].add_run(label); r.bold = True; r.font.size = Pt(8.5)
            b.text = ""
            hp = b.paragraphs[0]; hr = hp.add_run(hint); hr.font.size = Pt(6.5)
            hr.font.color.rgb = RGBColor.from_string(INK_SOFT.lstrip("#"))
            for _ in range(max(1, int(h / 5))):
                b.add_paragraph("")
        para("", 2, space_after=3)

    section("1 · Bemor")
    rows([PATIENT_FIELDS[0], ("Jinsi", "☐ Erkak    ☐ Ayol", 5)] + PATIENT_FIELDS[1:])
    section("2 · Klinik ma’lumot (davolovchi shifokor to‘ldiradi)")
    rows([CLINICAL_FIELDS[0], CLINICAL_FIELDS[1],
          ("Amaliyot turi", "☐ Eksizion  ☐ Insizion  ☐ Panch  ☐ Shave  ☐ Kyuretaj  ☐ Boshqa", 5),
          ("Amaliyot sanasi", "KK.OO.YYYY", 5), CLINICAL_FIELDS[3],
          ("Shoshilinchlik", "☐ Oddiy    ☐ STAT (shoshilinch)", 5),
          ("Ilova", "☐ Klinik surat (toshma)   ☐ Dermatoskopiya   ☐ Avvalgi gistologiya", 5),
          ("Davolovchi shifokor", "F.I.Sh., telefon, imzo", 6),
          ("Muassasa", "klinika / poliklinika nomi, shahar", 6)])
    section("3 · Patomorfologik tekshiruv (laboratoriya to‘ldiradi)")
    rows([("Qabul qilingan sana", "KK.OO.YYYY · qabul qilgan xodim", 6)] + LAB_FIELDS)
    para("Yo‘llanma va klinik surat DermaPATH ga birga yuklanadi: suratni tekis, yorug‘ joyda, butun varaq kadrda bo‘lib oling.",
         6.8, False, INK_SOFT, space_after=0)
    para("Yakuniy tashxis faqat litsenziyali patolog tomonidan qo‘yiladi. " + CLINIC + " · " + PRODUCT,
         6.8, False, INK_SOFT)
    doc.save(str(path))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    pdf, docx = OUT / "yollanma.pdf", OUT / "yollanma.docx"
    build_pdf(pdf)
    build_docx(docx)
    for p in (pdf, docx):
        print(f"{p.relative_to(ROOT)}  {p.stat().st_size // 1024} KB")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
